import io
import json
import os
import re
from pathlib import Path

from flask import Flask, abort, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from update_results import COLUMNS, MEDAL_FILL

DATA_DIR = Path(__file__).parent / "data"
MEDALS = ("platinum", "gold", "silver", "bronze")

app = Flask(__name__)


def load_store():
    path = DATA_DIR / "results.json"
    if not path.exists():
        return {"updated_at": "", "events": {}}
    return json.loads(path.read_text(encoding="utf-8"))


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/events")
def api_events():
    store = load_store()
    events = [
        {k: e[k] for k in ("event_id", "name", "url", "start_date", "end_date", "location")}
        for e in store["events"].values() if e["results"]
    ]
    events.sort(key=lambda e: e["start_date"], reverse=True)
    return jsonify(updated_at=store["updated_at"], events=events)


def _norm(text: str) -> str:
    """Lower case without spaces or punctuation: "Seraj E.H" and "seraj eh" match."""
    return re.sub(r"[^0-9a-z]+", "", text.lower())


def _profile(horse_url: str) -> str:
    """The horse's own page (all its shows) from an event's horse link."""
    m = re.search(r"/horses/[^/]+/\d+/?$", horse_url)
    return f"https://www.arabianessence.tv{m.group(0)}" if m else horse_url


def search(values) -> dict:
    """Results for the medal, event and horse names in a request's form or query."""
    medal = values.get("medal", "all").lower()
    if medal != "all" and medal not in MEDALS:
        raise ValueError(f"medal must be all or one of {', '.join(MEDALS)}")
    event_id = values.get("event", "")
    # Several names can be searched at once, separated by commas or new lines.
    raw = [t.strip() for t in re.split(r"[,;\n]+", values.get("q", "")) if _norm(t)]
    raw = list({_norm(t): t for t in reversed(raw)}.values())[::-1]  # drop repeats, keep order
    terms = [_norm(t) for t in raw]
    store = load_store()
    events = store["events"].values()
    if event_id:
        events = [e for e in events if e["event_id"] == event_id]
    rows = [r for e in events for r in e["results"]]

    not_found = []
    text = {}
    if terms:
        text = {id(r): (_norm(r["horse_name"]), _norm(r["championship"]) + "|" + _norm(r["event_name"]))
                for r in rows}

        def matches(r, t):
            horse, other = text[id(r)]
            return t in horse or t in other
        not_found = [orig for orig, t in zip(raw, terms) if not any(matches(r, t) for r in rows)]
        rows = [r for r in rows if any(matches(r, t) for t in terms)]

    counts = {m: 0 for m in MEDALS}
    horses: dict[str, dict] = {}
    for r in rows:
        counts[r["medal"]] += 1
        if terms and any(t in text[id(r)][0] for t in terms):
            h = horses.setdefault(_norm(r["horse_name"]), {
                "horse_name": r["horse_name"], "horse_url": _profile(r["horse_url"]),
                **{m: 0 for m in MEDALS}})
            h[r["medal"]] += 1

    if medal != "all":
        rows = [r for r in rows if r["medal"] == medal]
    rows.sort(key=lambda r: MEDALS.index(r["medal"]))
    rows.sort(key=lambda r: (r["start_date"], r["event_id"]), reverse=True)
    return {"updated_at": store["updated_at"], "terms": terms, "rows": rows, "counts": counts,
            "horses": sorted(horses.values(), key=lambda h: h["horse_name"]),
            "not_found": not_found}


@app.route("/api/results", methods=["GET", "POST"])
def api_results():
    try:
        found = search(request.values)
    except ValueError as exc:
        return jsonify(error=str(exc)), 400
    # A name search shows every win; browsing everything shows the latest ones.
    limit = 20000 if found["terms"] else 1000
    return jsonify(updated_at=found["updated_at"], total=len(found["rows"]),
                   results=found["rows"][:limit], counts=found["counts"],
                   horses=found["horses"], not_found=found["not_found"])


@app.route("/download/search.xlsx", methods=["GET", "POST"])
def download_search():
    """The current search (names, event and medal) as an Excel file."""
    try:
        found = search(request.values)
    except ValueError as exc:
        abort(400, str(exc))
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append([c[0] for c in COLUMNS])
    for r in found["rows"]:
        row = [", ".join(r[k]) if k == "videos" else r[k] for _, k in COLUMNS]
        row[0] = row[0].title()
        ws.append(row)
        ws.cell(ws.max_row, 1).fill = PatternFill("solid", fgColor=MEDAL_FILL[r["medal"]])
    sheets = [ws]
    if found["horses"]:
        hs = wb.create_sheet("Horses")
        hs.append(["Horse", *[m.title() for m in MEDALS], "Total", "Horse page"])
        for h in found["horses"]:
            hs.append([h["horse_name"], *[h[m] for m in MEDALS], sum(h[m] for m in MEDALS),
                       h["horse_url"]])
        sheets.append(hs)
    if found["not_found"]:
        ns = wb.create_sheet("Not found")
        ns.append(["Searched for, nothing found"])
        for t in found["not_found"]:
            ns.append([t])
        sheets.append(ns)
    for sheet in sheets:
        for cell in sheet[1]:
            cell.font = Font(bold=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for col in sheet.columns:
            width = max(len(str(c.value or "")) for c in col[:200])
            sheet.column_dimensions[col[0].column_letter].width = min(max(width + 2, 8), 60)
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, as_attachment=True, download_name="search-results.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/changes")
def api_changes():
    path = DATA_DIR / "changes.json"
    changes = json.loads(path.read_text(encoding="utf-8")) if path.exists() else []
    return jsonify(list(reversed(changes))[:200])


@app.get("/download/results.xlsx")
def download_excel():
    path = DATA_DIR / "results.xlsx"
    if not path.exists():
        abort(404)
    return send_file(path, as_attachment=True, download_name="championship-results.xlsx")


@app.get("/healthz")
def healthz():
    return "ok"


if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.environ.get("PORT", 5000)), debug=True)
